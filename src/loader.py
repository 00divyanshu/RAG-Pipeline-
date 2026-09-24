import logging
import csv
from pathlib import Path
from typing import List, Set
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: Set[str] = {".pdf", ".docx", ".csv", ".txt", ".md"}

def load_single_pdf(file_path: Path) -> List[Document]:
    """Loads a single PDF file and standardizes metadata."""
    path = Path(file_path)
    if not path.exists() or path.suffix.lower() != ".pdf":
        raise FileNotFoundError(f"Valid PDF file not found at: {path}")

    loader = PyPDFLoader(str(path))
    documents = loader.load()
    for doc in documents:
        doc.metadata["filename"] = path.name
        doc.metadata["file_type"] = "pdf"
        if "page" in doc.metadata:
            doc.metadata["page_number"] = doc.metadata["page"] + 1
    return documents

def load_single_docx(file_path: Path) -> List[Document]:
    """Loads a Microsoft Word (.docx) document, extracting text and tables."""
    path = Path(file_path)
    if not path.exists() or path.suffix.lower() != ".docx":
        raise FileNotFoundError(f"Valid DOCX file not found at: {path}")

    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required to load .docx files. Please install via pip install python-docx")

    doc = docx.Document(str(path))
    content_parts = []

    # 1. Paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            content_parts.append(text)

    # 2. Tables
    for t_idx, table in enumerate(doc.tables, 1):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_rows.append(" | ".join(row_cells))
        if table_rows:
            content_parts.append(f"\n[Table {t_idx}]\n" + "\n".join(table_rows) + "\n")

    full_text = "\n\n".join(content_parts)
    if not full_text:
        return []

    # Group into ~1500 char sections or pages
    chunk_size = 1500
    docs = []
    text_len = len(full_text)
    for idx, start in enumerate(range(0, text_len, chunk_size), 1):
        end = min(start + chunk_size, text_len)
        docs.append(Document(
            page_content=full_text[start:end],
            metadata={
                "source": str(path),
                "filename": path.name,
                "file_type": "docx",
                "page_number": idx,
            },
        ))
    return docs

def load_single_csv(file_path: Path) -> List[Document]:
    """Loads a CSV file and converts tabular rows into structured semantic documents."""
    path = Path(file_path)
    if not path.exists() or path.suffix.lower() != ".csv":
        raise FileNotFoundError(f"Valid CSV file not found at: {path}")

    docs: List[Document] = []
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            headers = next(reader, None)
            if not headers:
                return []

            batch_rows = []
            batch_size = 25
            batch_num = 1
            row_count = 0

            for row in reader:
                if not any(cell.strip() for cell in row):
                    continue
                row_count += 1
                row_desc = ", ".join([f"{h.strip()}: {c.strip()}" for h, c in zip(headers, row) if c.strip()])
                batch_rows.append(f"Row {row_count}: {row_desc}")

                if len(batch_rows) >= batch_size:
                    docs.append(Document(
                        page_content=f"Dataset: {path.name} (Columns: {', '.join(headers)})\n" + "\n".join(batch_rows),
                        metadata={
                            "source": str(path),
                            "filename": path.name,
                            "file_type": "csv",
                            "page_number": batch_num,
                        },
                    ))
                    batch_rows = []
                    batch_num += 1

            if batch_rows:
                docs.append(Document(
                    page_content=f"Dataset: {path.name} (Columns: {', '.join(headers)})\n" + "\n".join(batch_rows),
                    metadata={
                        "source": str(path),
                        "filename": path.name,
                        "file_type": "csv",
                        "page_number": batch_num,
                    },
                ))
    except Exception as e:
        logger.error(f"Failed parsing CSV {path}: {e}")
        raise e

    return docs

def load_single_text(file_path: Path) -> List[Document]:
    """Loads a plain text (.txt) or markdown (.md) document."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found at: {path}")

    ext = path.suffix.lower().lstrip(".")
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = path.read_text(encoding="latin-1", errors="replace")

    if not content.strip():
        return []

    # Partition large text files into 1500 char sections with section numbers
    chunk_size = 1500
    docs = []
    text_len = len(content)
    for idx, start in enumerate(range(0, text_len, chunk_size), 1):
        end = min(start + chunk_size, text_len)
        docs.append(Document(
            page_content=content[start:end],
            metadata={
                "source": str(path),
                "filename": path.name,
                "file_type": ext or "text",
                "page_number": idx,
            },
        ))
    return docs

def load_single_document(file_path: Path) -> List[Document]:
    """Unified document loader routing to the appropriate parser based on file extension."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return load_single_pdf(path)
    elif ext == ".docx":
        return load_single_docx(path)
    elif ext == ".csv":
        return load_single_csv(path)
    elif ext in [".txt", ".md"]:
        return load_single_text(path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}")

def load_documents_from_directory(directory_path: Path) -> List[Document]:
    """
    Loads all supported documents (PDF, DOCX, CSV, TXT, MD) from the specified directory.
    Attaches standardized metadata (filename, page/section number, file_type) to each document.
    """
    directory = Path(directory_path)
    if not directory.exists():
        directory.mkdir(parents=True, exist_ok=True)
        logger.warning(f"Created empty documents directory at {directory}")
        return []

    all_files = [f for f in directory.iterdir() if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS]
    if not all_files:
        logger.warning(f"No supported document files found in {directory}")
        return []

    logger.info(f"Found {len(all_files)} supported file(s) in {directory}: {[f.name for f in all_files]}")
    documents: List[Document] = []

    for file_path in sorted(all_files):
        try:
            doc_chunks = load_single_document(file_path)
            documents.extend(doc_chunks)
        except Exception as e:
            logger.error(f"Error loading {file_path.name}: {e}")

    logger.info(f"Loaded {len(documents)} document section(s) across {len(all_files)} file(s).")
    return documents
